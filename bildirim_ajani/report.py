"""Aylık bildirim raporu: numaralandırma, istatistik, Excel çıktısı ve Rapor Ajanı (yönetici özeti)."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .llm import LLMHatasi, LLMIstemci

KUMULE_KOLONLARI = [
    ("bildirim_no", "Bildirim No", 12),
    ("tarih", "Tarih", 12),
    ("gorusulen_mt", "Görüşülen MT", 16),
    ("bildirim_turu", "Bildirim Türü", 16),
    ("bildirim_sahibi", "Bildirim Sahibi", 18),
    ("telefon", "Telefon", 14),
    ("bildirim_yontemi", "Bildirim Yöntemi", 12),
    ("muhatap_kurum", "Muhatap Kurum", 16),
    ("fon_kodu", "Fon Kodu", 12),
    ("bildirimin_konusu", "Bildirimin Konusu", 28),
    ("bildirim_detayi", "Bildirim Detayı", 90),
    ("islem_ozeti", "İşlem Özeti", 24),
    ("cevaplama_tarihi", "Cevaplama Tarihi", 14),
    ("cevap", "Cevap", 10),
    ("cevaplama_suresi", "Cevaplama Süresi", 10),
    ("ai_duzeltmesi", "AI Düzeltmesi", 10),
    ("kaynak", "Kaynak", 8),
    ("analiz_guveni", "Analiz Güveni", 10),
    ("insan_onayi_gerekli", "İnsan Onayı Gerekli", 12),
    ("dosya_adi", "Kayıt Dosyası", 40),
]


def numaralandir(satirlar: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Satırları tarih ve kayıt saatine göre sıralayıp CM-0001'den başlayarak numaralar."""
    sirali = sorted(satirlar, key=lambda s: (s.get("tarih") or "9999", s.get("dosya_adi") or ""))
    for i, satir in enumerate(sirali, start=1):
        satir["bildirim_no"] = f"CM-{i:04d}"
    return sirali


def istatistik(satirlar: list[dict[str, Any]]) -> dict[str, Any]:
    """Raporun deterministik özet sayıları (Rapor Ajanı'na kaynak olarak verilir)."""
    fonlar: Counter[str] = Counter()
    for s in satirlar:
        for kod in (s.get("fon_kodu") or "-").split(","):
            if kod.strip() and kod.strip() != "-":
                fonlar[kod.strip()] += 1
    sikayet = [s for s in satirlar if s.get("bildirim_turu") == "Şikayet"]
    return {
        "toplam_bildirim": len(satirlar),
        "tur_dagilimi": dict(Counter(s.get("bildirim_turu", "") for s in satirlar).most_common()),
        "konu_dagilimi": dict(Counter(s.get("bildirimin_konusu", "") for s in satirlar).most_common()),
        "muhatap_dagilimi": dict(Counter(s.get("muhatap_kurum", "") for s in satirlar).most_common()),
        "sikayet_konulari": dict(Counter(s.get("bildirimin_konusu", "") for s in sikayet).most_common()),
        "en_cok_gecen_fonlar": dict(fonlar.most_common(8)),
        "temsilci_dagilimi": dict(Counter(s.get("gorusulen_mt", "") or "bilinmiyor" for s in satirlar).most_common()),
        "insan_onayi_bekleyen": sum(1 for s in satirlar if s.get("insan_onayi_gerekli")),
        "sikayet_orani_yuzde": round(100 * len(sikayet) / len(satirlar), 1) if satirlar else 0.0,
    }


def _sayi_normalize(n: str) -> str:
    n = n.replace(",", ".")
    return n.rstrip("0").rstrip(".") if "." in n else n


def _sayilar(nesne: Any) -> set[str]:
    bulunan: set[str] = set()
    if isinstance(nesne, dict):
        for k, v in nesne.items():
            bulunan |= _sayilar(k) | _sayilar(v)
    elif isinstance(nesne, (list, tuple)):
        for v in nesne:
            bulunan |= _sayilar(v)
    elif isinstance(nesne, (int, float)):
        bulunan.add(_sayi_normalize(str(nesne)))
    elif isinstance(nesne, str):
        bulunan |= {_sayi_normalize(n) for n in re.findall(r"\d+(?:[.,]\d+)?", nesne)}
    return bulunan


def yonetici_ozeti_yaz(llm: LLMIstemci, stats: dict[str, Any], satirlar: list[dict[str, Any]]) -> dict[str, Any]:
    """Rapor Ajanı: sayıları deterministik istatistikten alarak yönetici özeti yazar.

    Özetteki her sayı istatistikte var mı kontrol edilir; yoksa uyarı döner (uydurma sayı bekçisi).
    """
    sikayet_ornekleri = [s.get("bildirim_detayi", "")[:400] for s in satirlar if s.get("bildirim_turu") == "Şikayet"][:8]
    mesajlar = [
        {
            "role": "system",
            "content": "Sen bir portföy yönetim şirketinin müşteri deneyimi ve uyum raporlarını hazırlayan analistsin. "
            "Sana aylık müşteri bildirimlerinin istatistiği ve örnek şikayet özetleri verilir. Yönetim kuruluna sunulacak "
            "kısa bir özet yaz. SADECE verilen sayıları kullan, yeni sayı/yüzde üretme. Türkçe, resmi dil. "
            'Çıktı: {"ozet": "2 paragraf", "one_cikanlar": ["3-4 madde"], "oneriler": ["2-3 somut aksiyon önerisi"]}',
        },
        {
            "role": "user",
            "content": f"## İstatistik\n{json.dumps(stats, ensure_ascii=False)}\n\n## Örnek şikayet özetleri\n"
            + "\n".join(f"- {o}" for o in sikayet_ornekleri),
        },
    ]
    try:
        cikti = llm.json_iste(mesajlar, sicaklik=0.2)
    except (LLMHatasi, ValueError) as hata:
        return {"ozet": "", "one_cikanlar": [], "oneriler": [], "uyarilar": [f"Yönetici özeti üretilemedi: {hata}"]}
    metin = " ".join([str(cikti.get("ozet", ""))] + [str(x) for x in cikti.get("one_cikanlar", [])] + [str(x) for x in cikti.get("oneriler", [])])
    izinli = _sayilar(stats) | {"1", "2", "3", "4", "5"}  # madde numaralandırması için küçük sayılar serbest
    uydurma = sorted({n for n in re.findall(r"\d+(?:[.,]\d+)?", metin) if _sayi_normalize(n) not in izinli})
    return {
        "ozet": str(cikti.get("ozet", "")),
        "one_cikanlar": [str(x) for x in cikti.get("one_cikanlar", [])],
        "oneriler": [str(x) for x in cikti.get("oneriler", [])],
        "uyarilar": [f"Özette istatistikte olmayan sayı(lar) var, kontrol edin: {', '.join(uydurma)}"] if uydurma else [],
    }


def excel_yaz(satirlar: list[dict[str, Any]], yol: Path, stats: dict[str, Any] | None = None, ozet: dict[str, Any] | None = None) -> Path:
    """Kümüle (satır bazlı) + Özet + Yönetici Özeti sayfalarıyla Excel raporu yazar."""
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kümüle"
    baslik_dolgu = PatternFill("solid", fgColor="1F4E79")
    onay_dolgu = PatternFill("solid", fgColor="FCE4D6")
    for i, (_, baslik, genislik) in enumerate(KUMULE_KOLONLARI, start=1):
        hucre = ws.cell(row=1, column=i, value=baslik)
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = baslik_dolgu
        hucre.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = genislik
    for r, satir in enumerate(satirlar, start=2):
        for c, (anahtar, _, _) in enumerate(KUMULE_KOLONLARI, start=1):
            deger = satir.get(anahtar, "")
            if anahtar == "insan_onayi_gerekli":
                deger = "Evet" if deger else "Hayır"
            hucre = ws.cell(row=r, column=c, value=deger)
            hucre.alignment = Alignment(wrap_text=anahtar == "bildirim_detayi", vertical="top")
            if satir.get("insan_onayi_gerekli"):
                hucre.fill = onay_dolgu
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    stats = stats or istatistik(satirlar)
    oz = wb.create_sheet("Özet")
    satir_no = 1
    for baslik, tablo in [
        ("Bildirim Türü", stats["tur_dagilimi"]),
        ("Bildirimin Konusu", stats["konu_dagilimi"]),
        ("Muhatap Kurum", stats["muhatap_dagilimi"]),
        ("Şikayet Konuları", stats["sikayet_konulari"]),
        ("En Çok Geçen Fonlar", stats["en_cok_gecen_fonlar"]),
        ("Temsilci", stats["temsilci_dagilimi"]),
    ]:
        oz.cell(row=satir_no, column=1, value=baslik).font = Font(bold=True)
        oz.cell(row=satir_no, column=2, value="Adet").font = Font(bold=True)
        satir_no += 1
        for anahtar, adet in tablo.items():
            oz.cell(row=satir_no, column=1, value=anahtar)
            oz.cell(row=satir_no, column=2, value=adet)
            satir_no += 1
        satir_no += 1
    oz.cell(row=satir_no, column=1, value="Toplam bildirim").font = Font(bold=True)
    oz.cell(row=satir_no, column=2, value=stats["toplam_bildirim"])
    oz.cell(row=satir_no + 1, column=1, value="İnsan onayı bekleyen").font = Font(bold=True)
    oz.cell(row=satir_no + 1, column=2, value=stats["insan_onayi_bekleyen"])
    oz.column_dimensions["A"].width = 42

    if ozet and ozet.get("ozet"):
        yo = wb.create_sheet("Yönetici Özeti")
        yo.column_dimensions["A"].width = 120
        satirlar_yo = [("Özet", True), (ozet["ozet"], False), ("", False), ("Öne çıkanlar", True)]
        satirlar_yo += [(f"• {m}", False) for m in ozet.get("one_cikanlar", [])]
        satirlar_yo += [("", False), ("Öneriler", True)] + [(f"• {m}", False) for m in ozet.get("oneriler", [])]
        for i, (metin, kalin) in enumerate(satirlar_yo, start=1):
            hucre = yo.cell(row=i, column=1, value=metin)
            hucre.font = Font(bold=kalin)
            hucre.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(yol)
    return yol
