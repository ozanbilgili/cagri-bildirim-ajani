import json

from openpyxl import load_workbook

from bildirim_ajani.evaluation import alan_dogru_mu
from bildirim_ajani.rag import KuralDeposu
from bildirim_ajani.report import excel_yaz, istatistik, numaralandir, yonetici_ozeti_yaz
from bildirim_ajani.tools import AracBaglami, AracKutusu

from .conftest import ORNEK_TRANSKRIPT, SahteLLM


def _kutu(bilgi, hafiza, adlar):
    return AracKutusu(AracBaglami(bilgi, KuralDeposu(bilgi, None), hafiza, ORNEK_TRANSKRIPT), adlar)


def test_arac_hatalari_ajana_metin_olarak_doner(bilgi, hafiza):
    kutu = _kutu(bilgi, hafiza, ["fon_ara"])
    assert "hata" in json.loads(kutu.calistir("olmayan_arac", {}))
    assert "Eksik parametre" in json.loads(kutu.calistir("fon_ara", {}))["hata"]
    assert json.loads(kutu.calistir("fon_ara", {"sorgu": "o re te"}))[0]["kod"] == "ORT"


def test_kod_tarama_araci(bilgi, hafiza):
    sonuc = json.loads(_kutu(bilgi, hafiza, ["transkriptte_kod_tara"]).calistir("transkriptte_kod_tara", {}))
    assert sonuc[0]["kod"] == "ORT"


def test_numaralandirma_ve_istatistik():
    satirlar = numaralandir([
        {"tarih": "2026-09-05", "bildirim_turu": "Şikayet", "bildirimin_konusu": "Diğer", "muhatap_kurum": "Örnek Portföy", "fon_kodu": "ORT, OAL"},
        {"tarih": "2026-09-01", "bildirim_turu": "Bilgi Talebi", "bildirimin_konusu": "Diğer", "muhatap_kurum": "Örnek Portföy", "fon_kodu": "-"},
    ])
    assert [s["bildirim_no"] for s in satirlar] == ["CM-0001", "CM-0002"] and satirlar[0]["tarih"] == "2026-09-01"
    stats = istatistik(satirlar)
    assert stats["toplam_bildirim"] == 2 and stats["sikayet_orani_yuzde"] == 50.0
    assert stats["en_cok_gecen_fonlar"] == {"ORT": 1, "OAL": 1}


def test_excel_raporu(tmp_path):
    satirlar = numaralandir([{"tarih": "2026-09-01", "bildirim_turu": "Şikayet", "bildirimin_konusu": "Diğer", "muhatap_kurum": "Örnek Portföy", "fon_kodu": "-", "insan_onayi_gerekli": True}])
    yol = excel_yaz(satirlar, tmp_path / "r.xlsx", ozet={"ozet": "Özet", "one_cikanlar": ["a"], "oneriler": ["b"]})
    wb = load_workbook(yol)
    assert wb.sheetnames == ["Kümüle", "Özet", "Yönetici Özeti"]
    assert wb["Kümüle"]["A2"].value == "CM-0001"


def test_yonetici_ozeti_uydurma_sayi_bekcisi():
    llm = SahteLLM()
    llm.ozel["rapor"] = lambda m: {"ozet": "Bu ay 2 bildirim alındı, şikayet oranı %50 oldu; geçen aya göre 37 artış var.", "one_cikanlar": [], "oneriler": []}
    stats = istatistik([{"bildirim_turu": "Şikayet"}, {"bildirim_turu": "Bilgi Talebi"}])
    sonuc = yonetici_ozeti_yaz(llm, stats, [])
    assert sonuc["uyarilar"] and "37" in sonuc["uyarilar"][0] and "50" not in sonuc["uyarilar"][0]


def test_hafiza_duzeltme_ve_benzer_arama(orkestrator):
    sonuc = orkestrator.metinden_analiz(ORNEK_TRANSKRIPT, "26.09.03_101522_12_SELIN_ER_05550000123_08500000000_1.wav")
    orkestrator.hafiza.duzeltme_kaydet(sonuc.cagri_id, "fon_kodu", "ORT", "ORT, OAL", "altın fonu da konuşuldu")
    guncel = orkestrator.hafiza.cagri_getir(sonuc.cagri_id)
    assert guncel.satir["fon_kodu"] == "ORT, OAL" and guncel.alanlar["fon_kodu"].kaynak == "iş birimi düzeltmesi"
    benzer = orkestrator.hafiza.benzer_duzeltmeler("Kuzey Bank ÖRT fonu satamıyorum", ("fon_kodu",))
    assert benzer and benzer[0]["yeni"] == "ORT, OAL"


def test_degerlendirme_karsilastirma():
    assert alan_dogru_mu("fon_kodu", "ORT, OAL", "OAL,ORT")
    assert not alan_dogru_mu("fon_kodu", "-", "ORT")
    assert alan_dogru_mu("bildirim_sahibi", "Ayşe Kaya", "Ayşe")
    assert alan_dogru_mu("bildirim_sahibi", "", "")
    assert not alan_dogru_mu("bildirim_sahibi", "", "Ayşe")
